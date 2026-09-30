"""Excel workbooks of the CEDIS reports, sheet by sheet as the owner's
originals (2026-09-23): "modificaciones" (7 sheets) and "por hora" (7 sheets)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..motor import reporte_cedis as rc
from ..motor.reporte_cedis import BODEGON, EMPANADAS, HORAS, DatosCedis

VERDE_MARCA = "10564E"
ENCABEZADO_F = Font(bold=True, color="FFFFFF")
ENCABEZADO_R = PatternFill("solid", fgColor=VERDE_MARCA)
MONEDA, PCT, FECHA, ENTERO, DECIMAL = '"$"#,##0.00', "0.0%", "yyyy-mm-dd hh:mm:ss", "#,##0", "#,##0.00"


def _hoja(wb: Workbook, titulo: str, encabezado: list[str], filas: list[list], formatos: dict[int, str] | None = None,
          anchos: dict[int, int] | None = None):
    ws = wb.create_sheet(titulo[:31])
    ws.append(encabezado)
    for c in ws[1]:
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_R
    for fila in filas:
        ws.append([float(v) if hasattr(v, "as_tuple") else v for v in fila])  # Decimal -> number
    for col, fmt in (formatos or {}).items():
        for (celda,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            celda.number_format = fmt
    for col in range(1, len(encabezado) + 1):
        ws.column_dimensions[get_column_letter(col)].width = (anchos or {}).get(col, max(12, len(str(encabezado[col - 1])) + 2))
    ws.freeze_panes = "A2"
    return ws


def _notas(wb: Workbook, notas: list[str], datos: DatosCedis):
    ws = wb.create_sheet("Notas")
    ws.append(["Nota"])
    ws["A1"].font, ws["A1"].fill = ENCABEZADO_F, ENCABEZADO_R
    for n in notas:
        ws.append([n])
    ws.append([])
    ws.append(["Sucursal", "Inicio en Odoo"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for suc in datos.sucursales:
        ws.append([suc, datos.inicio_por_sucursal[suc]])
        ws.cell(ws.max_row, 2).number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 140
    for (celda,) in ws.iter_rows(min_row=2, max_col=1):
        celda.alignment = Alignment(wrap_text=True, vertical="top")


def _guardar(wb: Workbook) -> bytes:
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def modificaciones(datos: DatosCedis) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    filas = []
    resumen = rc.resumen(datos)
    for proveedor in (BODEGON, EMPANADAS):
        bloque = [f for f in resumen if f["proveedor"] == proveedor]
        for i, f in enumerate(bloque + ([rc.total_resumen(bloque)] if bloque else [])):
            filas.append([proveedor if i == 0 else None, f["sucursal"], f["ordenes"], f["modificadas"], f["pct"],
                          f["con_cantidad"], f["con_extras"], f["con_monto"], f["diferencia"]])
    _hoja(wb, "Resumen", ["Proveedor", "Sucursal", "Ordenes", "Modificadas", "% modificadas", "Con cambio de cantidad",
                          "Con lineas extra", "Con cambio de monto", "Diferencia neta $ (sin IVA)"], filas,
          {5: PCT, 9: MONEDA}, {1: 30, 2: 24})

    mods = [o for o in datos.ordenes if o.modificada]
    _hoja(wb, "Ordenes modificadas", ["Orden", "Sucursal", "Proveedor", "Confirmada (CDMX)", "Cambios de cantidad",
                                      "Lineas extra", "Cambios de monto", "Monto al confirmar", "Monto final",
                                      "Diferencia neta $", "Diferencia %", "Quedo en $0"],
          [[o.nombre, o.sucursal, o.proveedor, o.confirmada, len(o.cantidades), len(o.extras), len(o.montos),
            o.monto_al_confirmar, o.monto_final, o.diferencia,
            (o.diferencia / o.monto_al_confirmar) if o.monto_al_confirmar else None,
            "Si" if o.monto_final == 0 else None] for o in mods],
          {4: FECHA, 8: MONEDA, 9: MONEDA, 10: MONEDA, 11: PCT}, {2: 22, 3: 30, 4: 20})

    _hoja(wb, "Cambios de cantidad", ["Orden", "Sucursal", "Proveedor", "Confirmada (CDMX)", "Fecha cambio (CDMX)",
                                      "Horas despues de confirmar", "Usuario", "Producto", "Cantidad anterior",
                                      "Cantidad nueva", "Diferencia", "Tipo", "Precio unitario", "Impacto $ (sin IVA)"],
          [[o.nombre, o.sucursal, o.proveedor, o.confirmada, c.fecha, o.horas_despues(c.fecha), c.usuario, c.producto,
            c.anterior, c.nueva, c.diferencia, c.tipo, c.precio, c.impacto] for o in datos.ordenes for c in o.cantidades],
          {4: FECHA, 5: FECHA, 9: DECIMAL, 10: DECIMAL, 11: DECIMAL, 13: MONEDA, 14: MONEDA},
          {2: 22, 3: 30, 4: 20, 5: 20, 7: 34, 8: 36})

    _hoja(wb, "Lineas extra", ["Orden", "Sucursal", "Proveedor", "Confirmada (CDMX)", "Fecha cambio (CDMX)",
                               "Horas despues de confirmar", "Usuario", "Producto agregado", "Cantidad pedida",
                               "Cantidad recibida", "Precio unitario", "Subtotal $ (sin IVA)"],
          [[o.nombre, o.sucursal, o.proveedor, o.confirmada, e.fecha, o.horas_despues(e.fecha), e.usuario, e.producto,
            e.pedida, e.recibida, e.precio, e.subtotal] for o in datos.ordenes for e in o.extras],
          {4: FECHA, 5: FECHA, 9: DECIMAL, 10: DECIMAL, 11: MONEDA, 12: MONEDA}, {2: 22, 3: 30, 4: 20, 5: 20, 7: 34, 8: 36})

    _hoja(wb, "Cambios de monto", ["Orden", "Sucursal", "Proveedor", "Confirmada (CDMX)", "Fecha cambio (CDMX)",
                                   "Horas despues de confirmar", "Usuario", "Monto anterior (sin IVA)",
                                   "Monto nuevo (sin IVA)", "Diferencia $"],
          [[o.nombre, o.sucursal, o.proveedor, o.confirmada, m.fecha, o.horas_despues(m.fecha), m.usuario, m.anterior,
            m.nuevo, m.diferencia] for o in datos.ordenes for m in o.montos],
          {4: FECHA, 5: FECHA, 8: MONEDA, 9: MONEDA, 10: MONEDA}, {2: 22, 3: 30, 4: 20, 5: 20, 7: 34})

    _hoja(wb, "Productos mas cambiados", ["Proveedor", "Producto", "Tipo de cambio", "Cambios", "Aumentos",
                                          "Disminuciones", "Impacto $ (sin IVA)"],
          [[f["proveedor"], f["producto"], f["tipo"], f["cambios"], f["aumentos"], f["disminuciones"], f["impacto"]]
           for f in rc.productos_mas_cambiados(datos)], {7: MONEDA}, {1: 30, 2: 40, 3: 20})
    _notas(wb, rc.notas(datos, "modificaciones"), datos)
    return _guardar(wb)


def _tabla_horas(tabla: dict[int, dict[str, int]], sucursales: list[str]) -> list[list]:
    filas = [[rc.etiqueta_hora(h)] + [tabla[h].get(s, 0) for s in sucursales] + [sum(tabla[h].values())] for h in HORAS]
    return filas + [["Total"] + [sum(tabla[h].get(s, 0) for h in HORAS) for s in sucursales]
                    + [sum(sum(tabla[h].values()) for h in HORAS)]]


def por_hora(datos: DatosCedis) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    sucs = datos.sucursales
    encabezado = ["Hora (CDMX)"] + sucs + ["Total"]
    _hoja(wb, "Ordenes por hora", encabezado, _tabla_horas(rc.por_hora(datos), sucs), anchos={1: 16})
    _hoja(wb, "% Modificaciones", ["Sucursal", "Ordenes", "Con modificaciones", "% con modificaciones", "Ordenes Bodegon",
                                   "Modif. Bodegon", "Ordenes Empanadas", "Modif. Empanadas", "% Bodegon", "% Empanadas"],
          [[f["sucursal"], f["ordenes"], f["modificadas"], f["pct"], f["ordenes_bodegon"], f["modif_bodegon"],
            f["ordenes_empanadas"], f["modif_empanadas"], f["pct_bodegon"], f["pct_empanadas"]]
           for f in (lambda fs: fs + ([rc.total_pct_modificaciones(fs)] if fs else []))(rc.pct_modificaciones(datos))],
          {4: PCT, 9: PCT, 10: PCT}, {1: 24})
    todas, mods = rc.por_hora(datos), rc.por_hora(datos, solo_modificadas=True)

    def pct(h, s=None):
        total = sum(todas[h].values()) if s is None else todas[h].get(s, 0)
        m = sum(mods[h].values()) if s is None else mods[h].get(s, 0)
        return m / total if total else None
    ws = _hoja(wb, "% Modif. por hora", encabezado,
               [[rc.etiqueta_hora(h)] + [pct(h, s) for s in sucs] + [pct(h)] for h in HORAS], anchos={1: 16})
    for fila in ws.iter_rows(min_row=2, min_col=2):
        for celda in fila:
            celda.number_format = PCT
    _hoja(wb, "Por hora - Bodegon", encabezado, _tabla_horas(rc.por_hora(datos, BODEGON), sucs), anchos={1: 16})
    _hoja(wb, "Por hora - Empanadas", encabezado, _tabla_horas(rc.por_hora(datos, EMPANADAS), sucs), anchos={1: 16})
    _hoja(wb, "Detalle", ["Orden", "Sucursal", "Proveedor", "Estado", "Creada (CDMX)", "Confirmada (CDMX)", "Hora",
                          "Subtotal actual", "Modificada", "Detalle"],
          [[o.nombre, o.sucursal, o.proveedor, o.estado, o.creada, o.confirmada, o.creada.hour, o.monto_final,
            o.modificada_subtotal, rc.detalle_cambios(o)] for o in datos.ordenes],
          {5: FECHA, 6: FECHA, 8: MONEDA}, {2: 22, 3: 30, 5: 20, 6: 20, 10: 90})
    _notas(wb, rc.notas(datos, "por_hora"), datos)
    return _guardar(wb)


def nombre_archivo(tipo: str, datos: DatosCedis) -> str:
    return f"OC_Bodegon_Empanadas_{tipo}_{datos.desde:%Y-%m-%d}_a_{datos.hasta:%Y-%m-%d}.xlsx"
