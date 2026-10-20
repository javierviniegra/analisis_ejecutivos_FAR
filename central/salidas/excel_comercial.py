"""Excel workbook of the commercial report (openpyxl).

Per report: Resumen (header, Lectura, indicators table with real numbers so
they can be worked on), the data behind each chart with a native Excel chart,
Detalle (per branch or per day/block, as the PDF's second page) and Notas
(coverage, branches left out, rules and thresholds).
With several reports in one workbook, each sheet name starts with the branch.
"""

import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..motor.comparativos import BAJA, IGUAL, NARANJA, ROJO, SUBE, VERDE
from ..motor import detalle
from ..motor.periodo import TipoPeriodo
from ..motor.reporte_comercial import TITULO, ReporteComercial, filas_visibles, nombre_base
from ..motor.tabla_comercial import DECIMAL, ENTERO, MONEDA, PORCENTAJE
from . import AVISO_AUTOMATICO, pie_excel

VERDE_MARCA = "10564E"
VERDE_CLARO = "A7C2BC"
COLOR_VAR = {VERDE: "1E7B4F", ROJO: "B3261E"}
GRIS = "8A8A8A"
FONDO_CAJA = "F2F5F4"
FONDO_SEMAFORO = {VERDE: "CDEBD9", NARANJA: "FBE0C3", ROJO: "F6CFCB"}  # light fills behind the value

FORMATO_VALOR = {MONEDA: '"$"#,##0.00', ENTERO: "#,##0", PORCENTAJE: "0.0%", DECIMAL: "#,##0.0"}
FORMATO_VAR_RELATIVA = '+0.0%;-0.0%;0.0%'
FORMATO_VAR_PUNTOS = '+0.0" pp";-0.0" pp";0.0" pp"'  # the value is already in points
SIMBOLO = {SUBE: "▲", BAJA: "▼", IGUAL: "="}

TITULO_F = Font(bold=True, size=14, color=VERDE_MARCA)
SECCION_F = Font(bold=True, color=VERDE_MARCA)
ENCABEZADO_F = Font(bold=True, color="FFFFFF")
ENCABEZADO_R = PatternFill("solid", fgColor=VERDE_MARCA)
PAR_R = PatternFill("solid", fgColor=FONDO_CAJA)
SUAVE_F = Font(color="555555", size=9)


def _hoja(wb: Workbook, r: ReporteComercial, nombre: str, varios: bool):
    titulo = f"{r.nombre[:18]} · {nombre}" if varios else nombre
    return wb.create_sheet(titulo[:31])


def _variacion(ws, fila: int, col: int, var, fmt: str):
    """Numeric change in one cell (so it can be used in formulas) and the
    arrow, coloured, in the next one; 's/c' / 's/cf' when not comparable."""
    celda, flecha = ws.cell(fila, col), ws.cell(fila, col + 1)
    if var.porcentaje is None:
        flecha.value = var.simbolo
        flecha.font = Font(color=GRIS)
        return
    celda.value = float(var.porcentaje * 100 if fmt == PORCENTAJE else var.porcentaje)
    celda.number_format = FORMATO_VAR_PUNTOS if fmt == PORCENTAJE else FORMATO_VAR_RELATIVA
    flecha.value = SIMBOLO.get(var.direccion, "")
    flecha.font = Font(color=COLOR_VAR.get(var.color, GRIS), bold=True)
    flecha.alignment = Alignment(horizontal="center")


def _resumen(ws, r: ReporteComercial):
    p = r.periodo
    nombre = r.nombre if r.nombre != "Consolidado" else f"Consolidado ({len(r.sucursales)} sucursales)"
    ws["A1"] = f"{getattr(r, 'titulo', TITULO)} · {nombre}"
    ws["A1"].font = TITULO_F
    ws["A2"] = f"{p.etiqueta()}  ({p.desde:%d/%m/%Y} al {p.hasta:%d/%m/%Y})"
    anio = p.mismo_periodo_anio_anterior()
    solo_anio = getattr(r, "solo_anio", False)
    if solo_anio:
        ws["A3"] = (f"Contra: {anio.etiqueta()}  ({anio.desde:%d/%m/%Y} al {anio.hasta:%d/%m/%Y})" if anio
                    else "Contra: sin periodo comparable del año anterior")
    else:
        comparado = p.anterior().etiqueta() if anio is None or anio == p.anterior() else f"{p.anterior().etiqueta()} y {anio.etiqueta()}"
        ws["A3"] = f"Comparado con: {comparado}"
    ws["A3"].font = SUAVE_F
    if r.nombre == "Consolidado":
        ws["A4"] = "Sucursales: " + ", ".join(r.sucursales)
        ws["A4"].font = SUAVE_F
    aviso = getattr(r, "aviso_comparables", None)
    if aviso:
        ws.merge_cells("A5:G5")
        ws["A5"] = aviso
        ws["A5"].font = Font(bold=True, color="7A4E0F")
        ws["A5"].fill = PatternFill("solid", fgColor="FDF3E1")
        ws["A5"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[5].height = 30

    fila = 6
    ws.cell(fila, 1, r.lectura.titulo).font = SECCION_F
    bloques = [(b.subtitulo, b.observaciones) for b in r.lectura.bloques] or [("", r.lectura.observaciones)]
    for subtitulo, observaciones in bloques:
        if subtitulo:
            fila += 1
            ws.cell(fila, 1, subtitulo).font = Font(bold=True, color=VERDE_MARCA, size=9)
        for obs in observaciones:
            fila += 1
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
            celda = ws.cell(fila, 1, "• " + obs.texto)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[fila].height = 15 * (1 + len(obs.texto) // 110)

    fila += 2
    if solo_anio:
        _tabla_anio(ws, r, fila)
        return
    for col, texto in enumerate(["Indicador", "Periodo", "Anterior", "Var.", "", "Año anterior", "Var.", ""], start=1):
        c = ws.cell(fila, col, texto)
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_R
    seccion, par = None, False
    for f in filas_visibles(r.filas):
        if f.indicador.seccion != seccion:
            seccion = f.indicador.seccion
            fila += 1
            ws.cell(fila, 1, seccion).font = SECCION_F
            par = False
        fila += 1
        fmt = f.indicador.formato
        ws.cell(fila, 1, "   " + f.etiqueta)
        for col, valor in ((2, f.actual), (3, f.anterior), (6, f.anio_anterior)):
            c = ws.cell(fila, col, float(valor) if valor is not None else "—")
            c.number_format = FORMATO_VALOR[fmt]
            c.alignment = Alignment(horizontal="right")
        ws.cell(fila, 2).font = Font(bold=True)
        if f.semaforo:  # target indicator: the period's value on its traffic-light colour
            ws.cell(fila, 2).fill = PatternFill("solid", fgColor=FONDO_SEMAFORO[f.semaforo])
        _variacion(ws, fila, 4, f.var_anterior, fmt)
        _variacion(ws, fila, 7, f.var_anio, fmt)
        if par:
            for col in range(1, 9):
                if not (col == 2 and f.semaforo):  # keep the traffic-light colour
                    ws.cell(fila, col).fill = PAR_R
        par = not par
    for col, ancho in zip("ABCDEFGH", (38, 18, 18, 10, 4, 18, 10, 4)):
        ws.column_dimensions[col].width = ancho


def _tabla_anio(ws, r: ReporteComercial, fila: int):
    """The commercial table against last year only (owner, 2026-10-14)."""
    for col, texto in enumerate(["Indicador", "Periodo", "Año anterior", "Var.", ""], start=1):
        c = ws.cell(fila, col, texto)
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_R
    seccion, par = None, False
    for f in filas_visibles(r.filas):
        if f.indicador.seccion != seccion:
            seccion = f.indicador.seccion
            fila += 1
            ws.cell(fila, 1, seccion).font = SECCION_F
            par = False
        fila += 1
        fmt = f.indicador.formato
        ws.cell(fila, 1, "   " + f.etiqueta)
        for col, valor in ((2, f.actual), (3, f.anio_anterior)):
            c = ws.cell(fila, col, float(valor) if valor is not None else "—")
            c.number_format = FORMATO_VALOR[fmt]
            c.alignment = Alignment(horizontal="right")
        ws.cell(fila, 2).font = Font(bold=True)
        if f.semaforo:
            ws.cell(fila, 2).fill = PatternFill("solid", fgColor=FONDO_SEMAFORO[f.semaforo])
        _variacion(ws, fila, 4, f.var_anio, fmt)
        if par:
            for col in range(1, 6):
                if not (col == 2 and f.semaforo):
                    ws.cell(fila, col).fill = PAR_R
        par = not par
    for col, ancho in zip("ABCDE", (38, 18, 18, 10, 4)):
        ws.column_dimensions[col].width = ancho


def _datos_grafica(ws, g):
    """The chart's data as a table plus a native column chart."""
    ws["A1"] = g.titulo
    ws["A1"].font = SECCION_F
    hay_base = any(v is not None for v in g.base)
    encabezado = ["", g.nombre_actual] + ([g.nombre_base] if hay_base else [])
    for col, texto in enumerate(encabezado, start=1):
        c = ws.cell(3, col, texto)
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_R
    for i, etiqueta in enumerate(g.etiquetas):
        ws.cell(4 + i, 1, etiqueta.replace("\n", " "))
        for col, serie in ((2, g.actual), (3, g.base)):
            if col == 3 and not hay_base:
                continue
            valor = serie[i]
            c = ws.cell(4 + i, col, float(valor) if valor is not None else None)
            c.number_format = FORMATO_VALOR[MONEDA if g.pesos else ENTERO]
    ultima = 3 + len(g.etiquetas)
    if g.nota:
        ws.cell(ultima + 2, 1, g.nota).font = SUAVE_F
    ws.column_dimensions["A"].width = 12
    for col in ("B", "C"):
        ws.column_dimensions[col].width = 24

    grafica = BarChart()
    grafica.type, grafica.grouping, grafica.overlap = "col", "clustered", -10
    grafica.title = g.titulo
    grafica.y_axis.numFmt = '"$"#,##0' if g.pesos else "#,##0"
    grafica.height, grafica.width = 8, 18
    columnas = 3 if hay_base else 2
    grafica.add_data(Reference(ws, min_col=2, max_col=columnas, min_row=3, max_row=ultima), titles_from_data=True)
    grafica.set_categories(Reference(ws, min_col=1, min_row=4, max_row=ultima))
    for serie, color in zip(grafica.series, (VERDE_MARCA, VERDE_CLARO)):
        serie.graphicalProperties.solidFill = color
        serie.graphicalProperties.line.noFill = True
    ws.add_chart(grafica, f"{get_column_letter(columnas + 2)}3")


def _encabezado(ws, fila: int, titulos: list[str]):
    for col, texto in enumerate(titulos, start=1):
        c = ws.cell(fila, col, texto)
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_R


def _numero(ws, fila: int, col: int, valor, fmt: str):
    c = ws.cell(fila, col, float(valor) if valor is not None else None)
    c.number_format = FORMATO_VALOR[fmt]


def _detalle(ws, r: ReporteComercial):
    """Same detail as the PDF's second page: per branch, or per day/block."""
    ws["A1"] = "Detalle por sucursal" if r.por_sucursal else "Detalle por periodo"
    ws["A1"].font = SECCION_F
    if r.por_sucursal:
        costos = r.incluir_costos
        _encabezado(ws, 3, ["Sucursal", "Venta bruta", "Var. año ant.", "", "Clientes", "Cheque prom."]
                    + (["Costo total", "Costo facturado"] if costos else []))
        filas, total = detalle.por_sucursal(r)
        for i, f in enumerate(filas + [total], start=4):
            ws.cell(i, 1, f.nombre).font = Font(bold=f is total)
            _numero(ws, i, 2, f.venta_bruta, MONEDA)
            _variacion(ws, i, 3, f.var_anterior, MONEDA)
            _numero(ws, i, 5, f.clientes, ENTERO)
            _numero(ws, i, 6, f.cheque_promedio, MONEDA)
            for col, pct, color in ((7, f.pct_costo_total, f.semaforo_total), (8, f.pct_costo, f.semaforo))[:2 * costos]:
                _numero(ws, i, col, pct, PORCENTAJE)
                if color:
                    ws.cell(i, col).fill = PatternFill("solid", fgColor=FONDO_SEMAFORO[color])
            if costos and f.costo_total_estimado:  # estimated: "34.4%*", still a number
                ws.cell(i, 7).number_format = ws.cell(i, 7).number_format + '"*"'
        if costos and total.costo_total_estimado:
            ws.cell(5 + len(filas), 1, "* costo total estimado (ver hoja Notas)").font = Font(italic=True)
        anchos = (32, 18, 10, 4, 12, 14, 14, 16)
    else:
        compara = detalle.compara_con_anterior(r)
        _encabezado(ws, 3, ["Periodo", "Venta bruta", "Tickets", "Clientes", "Cheque prom."]
                    + (["Año anterior", "Var.", ""] if compara else []))
        for i, f in enumerate(detalle.por_periodo(r), start=4):
            ws.cell(i, 1, f.etiqueta.replace(chr(10), " "))
            _numero(ws, i, 2, f.venta_bruta, MONEDA)
            _numero(ws, i, 3, f.tickets, ENTERO)
            _numero(ws, i, 4, f.clientes, ENTERO)
            _numero(ws, i, 5, f.cheque_promedio, MONEDA)
            if compara:
                _numero(ws, i, 6, f.venta_anterior, MONEDA)
                _variacion(ws, i, 7, f.var_anterior, MONEDA)
        anchos = (14, 18, 10, 10, 14, 18, 10, 4)
    for col, ancho in zip("ABCDEFGH", anchos):
        ws.column_dimensions[col].width = ancho


def _notas(ws, r: ReporteComercial):
    ws["A1"] = "Notas"
    ws["A1"].font = SECCION_F
    for i, nota in enumerate(r.notas, start=3):
        c = ws.cell(i, 1, "• " + nota)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 15 * (1 + len(nota) // 120)
    ws.cell(len(r.notas) + 4, 1, AVISO_AUTOMATICO).font = Font(italic=True, color="555555")
    ws.column_dimensions["A"].width = 130


def _tablajeria(ws, r: ReporteComercial):
    """Same tables as the PDF's butchery page."""
    from ..motor.tablajeria import LEYENDA_DIAS, LEYENDA_PORCIONADOS, LEYENDA_PRODUCTOS, LEYENDA_SUCURSALES, notas
    t = r.tablajeria
    ws["A1"] = f"Tablajería · {r.nombre} · {r.periodo.etiqueta()}"
    ws["A1"].font = SECCION_F
    fila = 3
    costos = t.incluir_costos
    PUNTOS = "+0.0;-0.0;0.0"

    def puntos(col, v):
        ws.cell(fila, col, float(v * 100) if v is not None else None).number_format = PUNTOS

    def leyenda(texto):
        nonlocal fila
        ws.cell(fila - 1, 1, texto).font = Font(italic=True, color="555555", size=9)
        fila += 1

    def merma(titulo, filas, primera, referencia, vs):
        """`referencia`: (header, value of the row, its difference) after % merma, or None;
        `vs`: (header, difference) after the last-year columns, or None."""
        nonlocal fila
        ws.cell(fila, 1, titulo).font = Font(bold=True, color=VERDE_MARCA)
        fila += 1
        enc = [primera, "Kg base", "Merma kg", "% merma"]
        if referencia:
            enc += [referencia[0], f"Vs {referencia[0].lower()} (pp)"]
        enc += ["% año ant.", "Var. (pp)", ""] + ([vs[0]] if vs else []) + (["$ merma"] if costos else [])
        _encabezado(ws, fila, enc)
        for f in filas:
            fila += 1
            a, col = f.actual, 5
            ws.cell(fila, 1, f.etiqueta).font = Font(bold=f.total)
            _numero(ws, fila, 2, a.base, DECIMAL)
            _numero(ws, fila, 3, a.merma, DECIMAL)
            _numero(ws, fila, 4, f.indice, PORCENTAJE)
            if referencia:
                valor, dif = referencia[1](f)
                _numero(ws, fila, 5, valor, PORCENTAJE)
                puntos(6, dif)
                col = 7
            _numero(ws, fila, col, f.indice_anio, PORCENTAJE)
            _variacion(ws, fila, col + 1, f.var_anio, PORCENTAJE)
            col += 3
            if vs:
                puntos(col, vs[1](f))
                col += 1
            if costos:
                _numero(ws, fila, col, a.costo_merma, MONEDA)
        fila += 2

    if t.consolidado:
        merma("Merma por sucursal (de mayor a menor)", t.filas + ([t.total] if t.total else []), "Sucursal",
              ("% esperado", lambda f: (f.esperado, f.dif_esperado)), ("Vs grupo (pp)", lambda f: f.dif_grupo))
        leyenda(LEYENDA_SUCURSALES)
    else:
        merma("Por producto base (grupo = el mismo producto en todas las sucursales)",
              t.filas + ([t.total] if t.total else []), "Producto",
              ("% grupo", lambda f: (f.indice_grupo, f.dif_grupo)), None)
        leyenda(LEYENDA_PRODUCTOS)
        if t.por_dia:
            merma("Por día de la semana", t.por_dia, "Día", None, ("Vs su prom. (pp)", lambda f: f.dif_grupo))
            leyenda(LEYENDA_DIAS)
    for n in notas(t):
        ws.cell(fila, 1, "• " + n).font = Font(italic=True, color="555555", size=9)
        fila += 1
    for col, ancho in zip("ABCDEFGHIJK", (30, 11, 11, 10, 11, 12, 11, 10, 4, 14, 13)):
        ws.column_dimensions[col].width = ancho


def _composicion(ws, r: ReporteComercial):
    """Same data as the PDF's composition page."""
    from ..motor.composicion import HORAS
    k = r.composicion
    ws["A1"] = f"Composición de la venta · {r.nombre} · {r.periodo.etiqueta()}"
    ws["A1"].font = SECCION_F
    ws.cell(3, 1, "Venta por horario (hora de cierre del ticket)").font = Font(bold=True, color=VERDE_MARCA)
    _encabezado(ws, 4, ["Horario", "Horas", "Venta bruta", "% venta", "Tickets", "Clientes", "Cheque", "Ticket",
                        "Var. año ant.", ""])
    fila = 4
    for f in k.franjas:
        fila += 1
        ws.cell(fila, 1, f.etiqueta)
        ws.cell(fila, 2, f.horario)
        for col, v, fmt in ((3, f.venta, MONEDA), (4, f.pct, PORCENTAJE), (5, f.tickets, ENTERO),
                            (6, f.clientes, ENTERO), (7, f.cheque, MONEDA), (8, f.ticket, MONEDA)):
            _numero(ws, fila, col, v, fmt)
        _variacion(ws, fila, 9, f.var_anio, MONEDA)
    fila += 2
    if k.dias:
        ws.cell(fila, 1, "Por día de la semana (promedio por día)").font = Font(bold=True, color=VERDE_MARCA)
        fila += 1
        _encabezado(ws, fila, ["Día", "Días", "Mañana", "Tarde", "Noche", "Total", "% tarde", f"Vs {k.vs_anterior}", "",
                               "Vs año ant.", ""])
        for f in k.dias:
            fila += 1
            ws.cell(fila, 1, f.dia)
            ws.cell(fila, 2, f.dias)
            for col, v in zip((3, 4, 5, 6), f.bandas + [f.total]):
                _numero(ws, fila, col, v, MONEDA)
            _numero(ws, fila, 7, f.pct_tarde, PORCENTAJE)
            _variacion(ws, fila, 8, f.var_anterior, MONEDA)
            _variacion(ws, fila, 10, f.var_anio, MONEDA)
        fila += 2
    if k.tardes:
        ws.cell(fila, 1, "Quién vende mejor en las tardes (13:00 a 18:00), en proporción a su venta").font = Font(
            bold=True, color=VERDE_MARCA)
        fila += 1
        _encabezado(ws, fila, ["Sucursal", "Venta tarde", "% tarde", f"Vs {k.vs_anterior} (pp)", "",
                               "Vs año ant. (pp)", "", "% mañana", "% noche"])
        for f in k.tardes:
            fila += 1
            ws.cell(fila, 1, f.sucursal)
            _numero(ws, fila, 2, f.venta_tarde, MONEDA)
            _numero(ws, fila, 3, f.pct_tarde, PORCENTAJE)
            _variacion(ws, fila, 4, f.var_anterior, PORCENTAJE)
            _variacion(ws, fila, 6, f.var_anio, PORCENTAJE)
            _numero(ws, fila, 8, f.pct_manana, PORCENTAJE)
            _numero(ws, fila, 9, f.pct_noche, PORCENTAJE)
        fila += 2
    for titulo, rebanadas in k.pasteles:
        ws.cell(fila, 1, titulo).font = Font(bold=True, color=VERDE_MARCA)
        fila += 1
        _encabezado(ws, fila, ["Concepto", "Monto", "%"])
        total = sum((v for _, v, _ in rebanadas), 0) or 1
        for nombre, v, _ in rebanadas:
            fila += 1
            ws.cell(fila, 1, nombre)
            _numero(ws, fila, 2, v, MONEDA)
            _numero(ws, fila, 3, v / total, PORCENTAJE)
        fila += 2
    for g in k.barras:
        ws.cell(fila, 1, g.titulo).font = Font(bold=True, color=VERDE_MARCA)
        fila += 1
        _encabezado(ws, fila, ["Hora de cierre" if g.etiquetas == [str(h) for h in HORAS] else "Día", "Periodo",
                               "Año anterior"])
        for e, a, b in zip(g.etiquetas, g.actual, g.base):
            fila += 1
            ws.cell(fila, 1, e)
            _numero(ws, fila, 2, a, MONEDA)
            _numero(ws, fila, 3, b, MONEDA)
        if g.nota:
            fila += 1
            ws.cell(fila, 1, g.nota).font = Font(italic=True, color="555555", size=9)
        fila += 2
    for col, ancho in zip("ABCDEFGHIJ", (34, 16, 16, 10, 10, 10, 12, 12, 12, 4)):
        ws.column_dimensions[col].width = ancho


def generar(reportes: list[ReporteComercial]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    varios = len(reportes) > 1
    for r in reportes:
        _resumen(_hoja(wb, r, "Resumen", varios), r)
        if r.periodo.tipo == TipoPeriodo.MES:
            nombres = ["Venta por día", "Venta mensual"]
        else:
            nombres = ["Venta vs año anterior", "Clientes vs año anterior"]
        for g, nombre in zip(r.graficas, nombres):
            _datos_grafica(_hoja(wb, r, nombre, varios), g)
        if getattr(r, "tablajeria", None) is not None and not r.tablajeria.vacia:
            _tablajeria(_hoja(wb, r, "Tablajería", varios), r)
        if getattr(r, "composicion", None) is not None and not r.composicion.vacia:
            _composicion(_hoja(wb, r, "Composición", varios), r)
        _detalle(_hoja(wb, r, "Detalle", varios), r)
        _notas(_hoja(wb, r, "Notas", varios), r)
    pie_excel(wb)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def nombre_archivo(reportes: list[ReporteComercial]) -> str:
    return nombre_base(reportes) + ".xlsx"
