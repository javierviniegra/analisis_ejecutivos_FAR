import io
from datetime import date, datetime
from decimal import Decimal as D

import openpyxl
from django.test import SimpleTestCase

from .motor import reporte_cedis as rc
from .motor.fuentes import odoo
from .salidas import excel_cedis


def _msg(id, fecha_utc, cuerpo="", autor="Katia Moreno"):
    return {"id": id, "date": fecha_utc, "author_id": [1, autor], "body": cuerpo}


def _orden(nombre="P10276", sucursal="Antenas", proveedor=rc.BODEGON, final="11471.1", creada=datetime(2026, 6, 1, 10, 11)):
    return rc.Orden(nombre, sucursal, proveedor, "purchase", creada, datetime(2026, 6, 1, 11, 15, 20), D(final))


LINEAS = {"Queso Gouda": dict(precio=D("154"), pedida=D("0"), recibida=D("0")),
          "Gouda Polaco Recorte": dict(precio=D("140"), pedida=D("5"), recibida=D("5"))}


class InterpretarTests(SimpleTestCase):
    """Real P10276 (Antenas -> El Bodegón), from Odoo's chatter."""

    def _p10276(self):
        mensajes = [
            _msg(1, "2026-06-01 16:49:37"),  # before confirming: ignored
            _msg(2, "2026-06-01 19:00:57", "<p>Se actualizó la cantidad ordenada.</p><p>Queso Gouda: Cantidad ordenada: "
                                           "5.0 -&gt; 0.0 Cantidad facturada: 0.0</p>"),
            _msg(3, "2026-06-01 19:00:57", "<p>Linea extra con Gouda Polaco Recorte</p>"),
            _msg(4, "2026-06-01 19:00:57"),
        ]
        montos = {1: (D("0"), D("7580")), 4: (D("13069"), D("11471.1"))}
        return rc.interpretar(_orden(), mensajes, montos, LINEAS, odoo.a_cdmx)

    def test_solo_despues_de_confirmar(self):
        o = self._p10276()
        self.assertEqual(len(o.montos), 1)  # the pre-confirmation amount change is ignored
        self.assertEqual((o.monto_al_confirmar, o.diferencia), (D("13069.0"), D("-1597.9")))

    def test_cambio_de_cantidad_y_linea_extra(self):
        o = self._p10276()
        c = o.cantidades[0]
        self.assertEqual((c.producto, c.anterior, c.nueva, c.tipo, c.impacto), ("Queso Gouda", D("5.0"), D("0.0"), rc.DISMINUCION, D("-770.00")))
        self.assertEqual(c.fecha, datetime(2026, 6, 1, 13, 0, 57))  # UTC-6
        self.assertEqual(o.horas_despues(c.fecha), 1.8)
        e = o.extras[0]
        self.assertEqual((e.producto, e.pedida, e.subtotal), ("Gouda Polaco Recorte", D("5"), D("700")))

    def test_dos_definiciones_de_modificada(self):
        solo_cantidad = rc.interpretar(_orden("P12527", final="2400"), [_msg(1, "2026-06-02 10:00:00",
                                        "Se actualizó la cantidad ordenada. Te Sencha: Cantidad ordenada: 20.0 -> 80.0")],
                                       {}, {}, odoo.a_cdmx)
        self.assertTrue(solo_cantidad.modificada)  # "modificaciones": any change
        self.assertFalse(solo_cantidad.modificada_subtotal)  # "por hora": amount or extra line only
        self.assertIsNone(solo_cantidad.cantidades[0].impacto)  # no line price found: no invented impact

    def test_hora_cdmx(self):
        self.assertEqual(odoo.a_cdmx("2026-06-01 16:11:12"), datetime(2026, 6, 1, 10, 11, 12))
        self.assertEqual(odoo.limite_utc(date(2026, 6, 1)), "2026-06-01 06:00:00")


class TablasTests(SimpleTestCase):
    def _datos(self):
        a = InterpretarTests()._p10276()
        b = _orden("P2", "Antenas", rc.EMPANADAS, "2900", datetime(2026, 6, 1, 10, 21))
        c = _orden("P3", "Puebla", rc.BODEGON, "500", datetime(2026, 6, 2, 14, 5))
        return rc.DatosCedis(date(2026, 6, 1), date(2026, 9, 22), [a, b, c], ["Antenas", "Puebla"],
                             {"Antenas": date(2026, 6, 1), "Puebla": date(2026, 6, 10)}, 45, datetime(2026, 9, 30, 12, 0))

    def test_resumen_y_por_hora(self):
        d = self._datos()
        r = {(f["proveedor"], f["sucursal"]): f for f in rc.resumen(d)}
        f = r[(rc.BODEGON, "Antenas")]
        self.assertEqual((f["ordenes"], f["modificadas"], f["con_cantidad"], f["con_extras"], f["con_monto"], f["diferencia"]),
                         (1, 1, 1, 1, 1, D("-1597.9")))
        horas = rc.por_hora(d)
        self.assertEqual((horas[10]["Antenas"], horas[14]["Puebla"]), (2, 1))
        self.assertEqual(rc.por_hora(d, rc.EMPANADAS)[10]["Antenas"], 1)
        self.assertEqual(rc.detalle_cambios(d.ordenes[0]), "Monto 13,069.00 -> 11,471.10 | Linea extra con Gouda Polaco Recorte")  # as the owner's Excel

    def test_libros_con_las_hojas_originales(self):
        d = self._datos()
        m = openpyxl.load_workbook(io.BytesIO(excel_cedis.modificaciones(d)))
        self.assertEqual(m.sheetnames, ["Resumen", "Ordenes modificadas", "Cambios de cantidad", "Lineas extra",
                                        "Cambios de monto", "Productos mas cambiados", "Notas"])
        self.assertEqual(m["Resumen"]["K2"].value, -1597.9)
        self.assertEqual(m["Resumen"]["I1"].value, "Modificadas por la sucursal")
        self.assertEqual([m["Resumen"].cell(i, 2).value for i in range(2, 7)], ["Antenas", "Puebla", "Total", "Antenas", "Total"])
        self.assertIn("Se excluyen 45 órdenes generadas automáticamente", m["Notas"]["A5"].value)
        h = openpyxl.load_workbook(io.BytesIO(excel_cedis.por_hora(d)))
        self.assertEqual(h.sheetnames, ["Ordenes por hora", "% Modificaciones", "% Modif. por hora", "Por hora - Bodegon",
                                        "Por hora - Empanadas", "Detalle", "Notas"])
        self.assertEqual(h["Ordenes por hora"]["A26"].value, "Total")
        self.assertEqual(h["Ordenes por hora"]["D26"].value, 3)  # 3 orders in total


class QuienYClientesTests(SimpleTestCase):
    def test_hecho_por_la_sucursal_o_el_cedis(self):
        mensajes = [_msg(2, "2026-06-01 19:00:57", "Se actualizó la cantidad ordenada. Queso Gouda: Cantidad ordenada: 5.0 -> 0.0"),
                    _msg(3, "2026-06-01 20:00:00", "Linea adicional con Gouda Polaco Recorte")]  # sales-order wording
        mensajes[1]["author_id"] = [99, "Lilia Arredondo"]
        o = rc.interpretar(_orden(), mensajes, {}, LINEAS, odoo.a_cdmx, lambda autor: rc.CEDIS if autor == 99 else rc.SUCURSAL)
        self.assertEqual((o.cantidades[0].hecho_por, o.extras[0].hecho_por), (rc.SUCURSAL, rc.CEDIS))
        self.assertTrue(o.modificada_por(rc.CEDIS) and o.modificada_por(rc.SUCURSAL))

    def test_la_direccion_de_entrega_manda_sobre_el_cliente(self):
        class Fila:
            def __init__(self, nombre=None, excluir=False):
                self.nombre_reporte, self.excluir = nombre, excluir
        clientes = {1: Fila("Aeropuerto"), 2: Fila("Taquería Viaducto"), 3: Fila(excluir=True), 4: Fila()}
        pedido = lambda cliente, entrega: {"partner_id": [cliente, "x"], "partner_shipping_id": [entrega, "y"] if entrega else False}  # noqa: E731
        self.assertEqual(rc._cliente_de(pedido(1, 2), clientes).nombre_reporte, "Taquería Viaducto")
        self.assertEqual(rc._cliente_de(pedido(1, 1), clientes).nombre_reporte, "Aeropuerto")
        self.assertTrue(rc._cliente_de(pedido(3, None), clientes).excluir)
        self.assertIsNone(rc._cliente_de(pedido(4, 5), clientes))  # unassigned: reported in the notes
