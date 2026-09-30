from datetime import date, timedelta
from decimal import Decimal as D

from django.test import SimpleTestCase

from .motor.metricas import Comparacion, Metricas
from .motor.periodo import Periodo
from .motor.reporte_comercial import _uno, filas_visibles
from .salidas import excel_comercial, pdf_comercial


def _m(periodo):
    dias = {periodo.desde + timedelta(i): D("1000") for i in range(periodo.dias)}
    return Metricas(periodo=periodo, venta_bruta=D("7000"), venta_neta=D("6034"), tickets=D("70"), clientes=D("150"),
                    venta_por_dia=dias, dias_con_cierre=periodo.dias, dias_con_detalle=periodo.dias,
                    canal={"salon": D("6000"), "llevar": D("1000")}, mix={"Alimentos": D("5000"), "Bebidas": D("2000")})


class PdfTests(SimpleTestCase):
    SEM = Periodo.semana(2026, 38)

    def _reporte(self, nombre="Puebla"):
        ant, anio = self.SEM.anterior(), self.SEM.mismo_periodo_anio_anterior()
        return _uno(nombre, [nombre], self.SEM, _m(self.SEM), Comparacion(_m(self.SEM), _m(ant)),
                    Comparacion(_m(self.SEM), _m(anio)), None)

    def test_genera_dos_paginas_por_reporte(self):
        pdf = pdf_comercial.generar([self._reporte(), self._reporte("Acoxpa")])
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertEqual(pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page\r"), 4)  # report + detail each

    def test_oculta_secciones_sin_datos(self):
        visibles = filas_visibles(self._reporte().filas)
        self.assertNotIn("Costo", {f.indicador.seccion for f in visibles})
        self.assertIn("Ventas", {f.indicador.seccion for f in visibles})

    def test_nombre_de_archivo(self):
        self.assertEqual(pdf_comercial.nombre_archivo([self._reporte("La Esquina Coyoacán")]),
                         "Reporte_Comercial_La_Esquina_Coyoacán_Semana_38_14_sep_20_sep_2026.pdf")
        self.assertEqual(pdf_comercial.nombre_archivo([self._reporte(), self._reporte("Acoxpa")]),
                         "Reporte_Comercial_2_sucursales_Semana_38_14_sep_20_sep_2026.pdf")

    def test_excel_hojas_y_valores(self):
        import io

        import openpyxl
        libro = openpyxl.load_workbook(io.BytesIO(excel_comercial.generar([self._reporte()])))
        self.assertEqual(libro.sheetnames, ["Resumen", "Vs periodo anterior", "Vs año anterior", "Detalle", "Notas"])
        resumen = libro["Resumen"]
        self.assertEqual(resumen["A1"].value, "Reporte comercial · Puebla")
        filas = {str(r[0].value).strip(): r for r in resumen.iter_rows() if r[0].value}
        self.assertEqual(filas["Venta bruta (con IVA)"][1].value, 7000)  # a real number, not text
        self.assertNotIn("Costo", filas)  # empty section left out
        self.assertEqual(len(libro["Vs periodo anterior"]._charts), 1)
        dos = openpyxl.load_workbook(io.BytesIO(excel_comercial.generar([self._reporte(), self._reporte("Acoxpa")])))
        self.assertIn("Acoxpa · Resumen", dos.sheetnames)
        self.assertEqual(len(dos.sheetnames), 10)
